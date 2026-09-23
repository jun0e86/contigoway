"""업로드 문서(txt/csv/docx/xlsx/pptx/pdf/hwpx/hwp)에서 텍스트를 추출한다.

- 모두 메모리에서만 처리하고 디스크에 저장하지 않는다.
- 외부에서 온 파일이므로 zip 폭탄/XML 폭탄/압축 폭탄을 방어한다.
"""
from __future__ import annotations

import io
import re
import zipfile
import zlib
from dataclasses import dataclass, field
from pathlib import Path


class ExtractError(Exception):
    """사용자에게 그대로 보여줘도 되는 추출 실패 메시지."""


@dataclass
class ExtractedDoc:
    filename: str
    text: str
    warnings: list = field(default_factory=list)


MAX_UNCOMPRESSED = 200 * 1024 * 1024  # zip 계열 파일의 압축 해제 총량 상한


def extract(filename: str, data: bytes) -> ExtractedDoc:
    ext = Path(filename).suffix.lower()
    handlers = {
        ".txt": _plain, ".md": _plain, ".csv": _plain,
        ".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".pptx": _pptx,
        ".pdf": _pdf, ".hwpx": _hwpx, ".hwp": _hwp,
    }
    if ext in (".doc", ".xls", ".ppt"):
        raise ExtractError(
            f"{filename}: 구형 Office 형식({ext})은 지원하지 않습니다. "
            f"{ext}x 로 다시 저장하거나 PDF로 변환해 올려주세요."
        )
    if ext not in handlers:
        raise ExtractError(f"{filename}: 지원하지 않는 형식입니다. (지원: txt, csv, docx, xlsx, pptx, pdf, hwpx, hwp)")
    try:
        text, warnings = handlers[ext](data)
    except ExtractError as e:
        raise ExtractError(f"{filename}: {e}") from None
    except Exception as e:  # 손상된 파일 등
        raise ExtractError(f"{filename}: 파일을 읽을 수 없습니다. ({type(e).__name__})") from None
    text = _normalize(text)
    if not text:
        warnings.append("추출된 텍스트가 없습니다. (이미지/스캔 문서일 수 있음)")
    return ExtractedDoc(filename=filename, text=text, warnings=warnings)


def _normalize(text: str) -> str:
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)  # 제어문자 제거
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _check_zip(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED:
            raise ExtractError("압축 해제 크기가 너무 큽니다.")


# ── 텍스트 ────────────────────────────────────────────────
def _plain(data: bytes):
    for enc in ("utf-8-sig", "cp949"):
        try:
            return data.decode(enc), []
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore"), ["인코딩을 확정하지 못해 일부 글자가 깨졌을 수 있습니다."]


# ── Word ─────────────────────────────────────────────────
def _docx(data: bytes):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_zip(data)
    doc = Document(io.BytesIO(data))
    lines = []
    for child in doc.element.body.iterchildren():  # 본문 순서 그대로 (문단·표 섞임 유지)
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, doc).text.strip()
            if t:
                lines.append(t)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                seen, cells = set(), []
                for c in row.cells:
                    if id(c._tc) in seen:  # 병합 셀 중복 제거
                        continue
                    seen.add(id(c._tc))
                    cells.append(c.text.strip().replace("\n", " "))
                if any(cells):
                    lines.append(" | ".join(cells))
    return "\n".join(lines), []


# ── Excel ────────────────────────────────────────────────
def _xlsx(data: bytes):
    import openpyxl

    _check_zip(data)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines, warnings = [], []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        lines.append(f"[시트: {ws.title}]")
        for n, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if n > 3000:
                warnings.append(f"시트 '{ws.title}'는 3,000행까지만 읽었습니다.")
                break
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                lines.append(" | ".join(cells).rstrip(" |"))
    wb.close()
    return "\n".join(lines), warnings


# ── PowerPoint ───────────────────────────────────────────
def _pptx(data: bytes):
    from pptx import Presentation

    _check_zip(data)
    prs = Presentation(io.BytesIO(data))
    lines = []

    def walk(shapes):
        for sh in shapes:
            if getattr(sh, "shapes", None) is not None and sh.shape_type == 6:  # 그룹
                walk(sh.shapes)
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                lines.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))

    for i, slide in enumerate(prs.slides, start=1):
        lines.append(f"[슬라이드 {i}]")
        walk(slide.shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                lines.append(f"(발표자 노트) {note}")
    return "\n".join(lines), []


# ── PDF ──────────────────────────────────────────────────
def _pdf(data: bytes):
    import pdfplumber

    lines, warnings = [], []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = pdf.pages
        if len(pages) > 200:
            warnings.append("200페이지까지만 읽었습니다.")
            pages = pages[:200]
        for i, page in enumerate(pages, start=1):
            t = (page.extract_text() or "").strip()
            if t:
                lines.append(f"[{i}쪽]\n{t}")
    text = "\n".join(lines)
    if len(text) < 50:
        warnings.append("텍스트를 거의 읽지 못했습니다. 스캔본이면 텍스트 PDF/Word로 다시 올려주세요. (OCR 미지원)")
    return text, warnings


# ── HWPX (zip + XML) ─────────────────────────────────────
def _hwpx(data: bytes):
    from defusedxml.ElementTree import iterparse

    _check_zip(data)
    lines = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        sections = sorted(
            (n for n in z.namelist() if re.fullmatch(r"Contents/section\d+\.xml", n)),
            key=lambda n: int(re.findall(r"\d+", n)[-1]),
        )
        if not sections:
            raise ExtractError("HWPX 본문(Contents/section*.xml)을 찾지 못했습니다.")
        for name in sections:
            buf, parts = [], []
            for event, el in iterparse(io.BytesIO(z.read(name)), events=("end",)):
                tag = el.tag.rsplit("}", 1)[-1]
                if tag == "t":
                    parts.append("".join(el.itertext()))
                elif tag == "p":  # 문단 끝
                    line = "".join(parts).strip()
                    parts = []
                    if line:
                        buf.append(line)
            lines.extend(buf)
    return "\n".join(lines), []


# ── HWP 5.x (OLE 바이너리) ───────────────────────────────
_HWPTAG_PARA_TEXT = 16 + 51


def _hwp_records(buf: bytes):
    pos, n = 0, len(buf)
    while pos + 4 <= n:
        header = int.from_bytes(buf[pos:pos + 4], "little")
        pos += 4
        tag, size = header & 0x3FF, (header >> 20) & 0xFFF
        if size == 0xFFF:
            if pos + 4 > n:
                break
            size = int.from_bytes(buf[pos:pos + 4], "little")
            pos += 4
        yield tag, buf[pos:pos + size]
        pos += size


def _hwp_para_text(rec: bytes) -> str:
    """PARA_TEXT 레코드(UTF-16LE)에서 제어문자를 걷어내고 글자만 뽑는다."""
    out, i, n = bytearray(), 0, len(rec)
    while i + 1 < n:
        ch = rec[i] | (rec[i + 1] << 8)
        if ch in (10, 13):                    # 줄바꿈/문단 끝
            out += b"\n\x00"
            i += 2
        elif ch in (0, 24, 25, 26, 27, 28, 29):  # 1글자짜리 제어문자
            i += 2
        elif ch in (30, 31):                  # 고정폭/묶음 공백
            out += b" \x00"
            i += 2
        elif ch < 32:                         # 인라인/확장 제어(총 16바이트 차지)
            if ch == 9:
                out += b" \x00"
            i += 16
        else:
            out += rec[i:i + 2]
            i += 2
    return out.decode("utf-16-le", errors="ignore")


def _hwp(data: bytes):
    import olefile

    if not olefile.isOleFile(io.BytesIO(data)):
        raise ExtractError("HWP 5.x 형식이 아닙니다. (HWPX이거나 손상된 파일일 수 있음)")
    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        header = ole.openstream("FileHeader").read()
        if not header.startswith(b"HWP Document File"):
            raise ExtractError("HWP 파일 헤더가 올바르지 않습니다.")
        flags = int.from_bytes(header[36:40], "little")
        compressed, encrypted, distribution = bool(flags & 1), bool(flags & 2), bool(flags & 4)
        if encrypted:
            raise ExtractError("암호가 설정된 HWP는 읽을 수 없습니다. 암호 해제 후 올려주세요.")
        if distribution:
            raise ExtractError("배포용 HWP는 읽을 수 없습니다. 일반 문서로 저장하거나 PDF/HWPX로 변환해 올려주세요.")

        sections = sorted(
            (e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText" and e[1].startswith("Section")),
            key=lambda e: int(re.findall(r"\d+", e[1])[-1]),
        )
        lines = []
        for entry in sections:
            raw = ole.openstream(entry).read()
            if compressed:
                raw = zlib.decompressobj(-15).decompress(raw, 100 * 1024 * 1024)
            for tag, rec in _hwp_records(raw):
                if tag == _HWPTAG_PARA_TEXT:
                    t = _hwp_para_text(rec).strip()
                    if t:
                        lines.append(t)
    finally:
        ole.close()
    warnings = ["HWP(구형)는 표·서식이 일부 누락될 수 있습니다. 가능하면 HWPX/PDF로 올려주세요."]
    return "\n".join(lines), warnings
