"""검토 결과를 Word(.docx)로 만든다. (A4 가로, 체크리스트 표 중심)"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .pipeline import ReviewResult

FONT = "맑은 고딕"
# 판정별 (배경, 글자색)
VERDICT_STYLE = {
    "양호": ("C6EFCE", "006100"),
    "주의": ("FFEB9C", "7F4F00"),
    "위반 소지": ("FFC7CE", "9C0006"),
    "확인 필요": ("E7E6E6", "404040"),
}
NAVY = "1F3864"

DISCLAIMER = [
    "본 문서는 AI가 국가법령정보센터에서 조회한 법령 조문과 제출된 문서·요청 내용을 바탕으로 작성한 1차 검토 자료이며, 법률 자문이 아닙니다.",
    "최종 판단 및 조치 전에 법무 담당자 또는 변호사의 검토를 받으시기 바랍니다.",
    "검토에는 조회된 법령 조문만 사용했으며, 고시·지침·판례·유권해석·별표 등은 반영되지 않았을 수 있습니다.",
    "‘양호’는 제출된 문서와 요청 내용에서 확인된 범위 안에서 위반 소지를 발견하지 못했다는 의미입니다.",
]


def _clean(text) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", "", str(text or ""))


def _set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _write(cell, text, size=9, bold=False, color=None, align=None):
    """셀에 텍스트를 쓴다. 줄바꿈은 문단으로 분리."""
    lines = _clean(text).split("\n") or [""]
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if align:
            p.alignment = align
        _set_font(p.add_run(line), size=size, bold=bold, color=color)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _set_widths(table, widths_cm):
    table.autofit = False   # w:tblLayout fixed (python-docx가 올바른 위치에 삽입)
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w / 2.54 * 1440)))


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    _set_font(p.add_run(_clean(text)), size=13, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", NAVY)):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    pPr.insert_element_before(
        bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
        "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
        "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )


def _add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _set_font(run, size=8, color="7F7F7F")
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def build_docx(r: ReviewResult) -> bytes:
    doc = Document()
    for z in doc.settings.element.findall(qn("w:zoom")):   # python-docx 기본 템플릿의 누락 속성 보정
        if z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(10)

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    _add_page_number(sec)

    # 제목
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    _set_font(p.add_run("의료 관련 법령 검토 결과"), size=20, bold=True, color=NAVY)
    p = doc.add_paragraph()
    _set_font(p.add_run(f"검토 일시 {r.reviewed_at}  |  AI 1차 검토 (법률 자문 아님)"), size=9, color="7F7F7F")

    # 1. 검토 개요
    _heading(doc, "1. 검토 개요")
    laws_txt = "\n".join(
        f"{l['name']}" + (f" (시행 {l['시행일자']})" if l.get("시행일자") else "") for l in r.laws
    )
    rows = [
        ("검토 요청", r.request_text or "(요청 문구 없음 — 문서 기준 검토)"),
        ("요약", r.request_summary),
        ("검토 대상 문서", "\n".join(r.files) if r.files else "(없음)"),
        ("조회한 법령", laws_txt),
    ]
    t = doc.add_table(rows=len(rows) + 1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        _shade(t.cell(i, 0), "F2F2F2")
        _write(t.cell(i, 0), k, bold=True)
        _write(t.cell(i, 1), v)
    bg, fg = VERDICT_STYLE[r.overall_verdict]
    _shade(t.cell(len(rows), 0), "F2F2F2")
    _write(t.cell(len(rows), 0), "종합 판정", bold=True)
    _shade(t.cell(len(rows), 1), bg)
    _write(t.cell(len(rows), 1), r.overall_verdict, size=11, bold=True, color=fg)
    _set_widths(t, [4.0, 21.7])

    # 2. 종합 의견
    _heading(doc, "2. 종합 의견")
    for line in _clean(r.overall_comment).split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _set_font(p.add_run(line), size=10)

    # 3. 항목별 체크리스트
    _heading(doc, "3. 항목별 검토 결과 (체크리스트)")
    headers = ["No", "확인 항목", "관련 조문", "판정", "검토 내용 및 문서상 근거", "조치사항"]
    widths = [1.0, 4.2, 3.6, 2.0, 8.2, 6.7]
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(tbl.rows[0].cells, headers):
        _shade(c, NAVY)
        _write(c, h, size=9, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    _repeat_header(tbl.rows[0])
    for n, it in enumerate(r.items, start=1):
        cells = tbl.add_row().cells
        refs = list(it["refs"]) + [f"{x} (원문 미확인)" for x in it["unverified_refs"]]
        bg, fg = VERDICT_STYLE[it["verdict"]]
        _write(cells[0], str(n), align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(cells[1], it["title"], bold=True)
        _write(cells[2], "\n".join(refs) or "-", size=8.5)
        _shade(cells[3], bg)
        _write(cells[3], it["verdict"], bold=True, color=fg, align=WD_ALIGN_PARAGRAPH.CENTER)
        body = it["finding"] + (f"\n▶ 문서상 근거: {it['evidence']}" if it["evidence"] else "")
        _write(cells[4], body)
        _write(cells[5], it["action"])
    if not r.items:
        cells = tbl.add_row().cells
        _write(cells[1], "검토 항목이 생성되지 않았습니다.")
    _set_widths(tbl, widths)

    # 4. 추가 확인 사항 / 제한 사항
    if r.missing_info or r.warnings:
        _heading(doc, "4. 추가 확인 필요 사항 및 검토 제한")
        for title, arr in (("추가로 필요한 자료·확인 사항", r.missing_info), ("검토 제한 사항", r.warnings)):
            if not arr:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            _set_font(p.add_run(title), size=10, bold=True)
            for x in arr:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
                _set_font(p.add_run(_clean(x)), size=9.5)

    # 부록: 조문 원문
    if r.articles:
        doc.add_page_break()
        _heading(doc, "부록. 검토에 인용된 조문 원문")
        p = doc.add_paragraph()
        _set_font(p.add_run("국가법령정보센터 조회 기준. 아래 원문을 근거로 판정했습니다."), size=9, color="7F7F7F")
        for a in r.articles:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            _set_font(p.add_run(f"{a.ref}({a.title})" if a.title else a.ref), size=10, bold=True)
            for line in _clean(a.text).split("\n"):
                q = doc.add_paragraph()
                q.paragraph_format.space_after = Pt(1)
                q.paragraph_format.left_indent = Cm(0.4)
                _set_font(q.add_run(line), size=9)

    # 유의사항
    _heading(doc, "유의사항")
    for line in DISCLAIMER:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        _set_font(p.add_run(line), size=9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
